from bs4 import BeautifulSoup
from selenium import webdriver
import selenium
from docx import Document

document = Document()
browser = webdriver.Chrome()
browser.delete_all_cookies()
site = 'https://sites.google.com/site/hackbookeasypeasy/'
thehtml = browser.get(site)
browser.find_element_by_id('sites-mature-confirm-btn').click()
#browser.find_element_by_partial_link_text('NEXT PAGE').click()

n = 0
while True:
    title = browser.find_element_by_id('sites-page-title')
    try:
        textp = browser.find_element_by_class_name('sites-tile-name-content-2')
    except:
        textp = browser.find_element_by_class_name('sites-tile-name-content-1')

    document.add_heading(title.text, 0)
    document.add_paragraph(textp.text)
    if 'Download' in title.text:
        break

    try:
        browser.find_element_by_partial_link_text('NEXT').click()
    except:
        browser.find_elements_by_class_name('sites-navigation-link')[n+1].click()

    n += 1

document.save('easypeasy.docx')


#article = soup.select('font', attrs={'size' : '4'})
#print(article)
# textl = []
# for art in article:
#     text = art.text.strip(' ')
#     textl.append(text)

#textl = list(dict.fromkeys(textl))
# nextp = soup.select('a')[2].get('href')
#print(soup.prettify)
#print(textl)
